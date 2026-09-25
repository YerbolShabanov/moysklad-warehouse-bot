"""Маршрутизация заказов и сводный сборочный лист (очередь, не живой опрос)."""
import asyncio
import io
import json
import sys
import tempfile
from datetime import datetime, timedelta
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from openpyxl import load_workbook

from msbot import routing
from msbot.batcher import PickingBatcher
from msbot.config import Settings
from msbot.excel import render_batch_xlsx
from msbot.moysklad import MoySkladError
from msbot.orders import OrderCard, OrderPosition, OrderService, SlotStock
from test_smoke import FakeClient

KASPI = {"salesChannel": {"name": "Kaspi магазин"}, "agent": {"name": "Kaspi магазин"}}
EXPRESS_ON = [{"name": "Экспресс доставка", "value": True}]


def test_routing():
    assert routing.route(KASPI, "Kaspi магазин", "Экспресс доставка") == routing.BATCH
    express = dict(KASPI, attributes=EXPRESS_ON)
    assert routing.route(express, "Kaspi магазин", "Экспресс доставка") == routing.EXPRESS
    # Снятая галочка приходит как отсутствие атрибута, а не как false
    off = dict(KASPI, attributes=[{"name": "Тип доставки", "value": "Kaspi Postomat"}])
    assert routing.route(off, "Kaspi магазин", "Экспресс доставка") == routing.BATCH
    assert routing.is_express(dict(KASPI, attributes=[{"name": "Экспресс доставка",
                                                       "value": False}]),
                              "Экспресс доставка") is False

    opt = {"salesChannel": {"name": "ОПТ"}, "agent": {"name": "Карабаликов ИП"}}
    assert routing.route(opt, "Kaspi магазин", "Экспресс доставка") == routing.INDIVIDUAL
    # Совпадения только по контрагенту достаточно
    by_agent = {"salesChannel": {"name": "Точка продаж"}, "agent": {"name": "Kaspi магазин"}}
    assert routing.route(by_agent, "Kaspi магазин", "Экспресс доставка") == routing.BATCH


def card(number, positions, state="Упаковка"):
    return OrderCard(id=f"id-{number}", number=number, state=state,
                     moment=datetime.now(), sales_channel="Kaspi магазин",
                     agent="Kaspi магазин", store="Склад", positions=positions)


def pos(name, barcode, row, place, level, qty=1):
    p = OrderPosition(name=name, code="", barcode=barcode, quantity=qty, uom="шт")
    p.slots = [SlotStock(slot_name=f"Ряд {row} место {place} этаж {level}", stock=9)]
    return p


def test_sheet():
    cards = [
        card("1001", [pos("Автокресло OXIS", "4690624238147", 10, 2, 1),
                      pos("Ходунки SMILEY", "0004690624101601", 3, 4, 1)]),
        card("1002", [pos("Беговел WALLY", "4690624197727", 3, 1, 1, qty=2)], state="Новый"),
    ]
    book = load_workbook(io.BytesIO(render_batch_xlsx(cards)))
    assert book.sheetnames == ["Маршрут сборки", "По заказам"], book.sheetnames

    route = book["Маршрут сборки"]
    head_row = next(r for r in range(1, route.max_row + 1)
                    if route.cell(row=r, column=1).value == "№")
    head = [route.cell(row=head_row, column=c).value for c in range(1, 7)]
    assert head == ["№", "Наименование товара", "Штрихкод", "Кол-во", "Ячейка", "Заказ"], head

    body = [[route.cell(row=r, column=c).value for c in range(1, 7)]
            for r in range(head_row + 1, route.max_row + 1)]
    # Маршрут не зависит от заказа: идём по рядам подряд
    assert [r[4] for r in body] == ["Р3-М1-Э1", "Р3-М4-Э1", "Р10-М2-Э1"], body
    # У каждой строки виден свой заказ
    assert [r[5] for r in body] == ["1002", "1001", "1001"], body
    # Штрихкод текстом — ведущий ноль на месте
    assert body[1][2] == "0004690624101601", body[1]

    orders = book["По заказам"]
    ohead = next(r for r in range(1, orders.max_row + 1)
                 if orders.cell(row=r, column=1).value == "Заказ")
    numbers = [orders.cell(row=r, column=1).value
               for r in range(ohead + 1, orders.max_row + 1)]
    # Номер печатается один раз на группу, дальше пусто (openpyxl отдаёт None)
    assert numbers == ["1001", None, "1002"], numbers


class BatchTestClient(FakeClient):
    """Заказы адресуются по номеру (`name`), как их находит `card_by_number`.

    Остатки по ячейкам фиксированные — если не сказано иное, card никогда
    не окажется `slots_failed`, если явно не попросить обратное.
    """

    def __init__(self, orders):
        super().__init__()
        self.orders = {o["name"]: dict(o) for o in orders}
        self.stock_fails = 0  # сколько раз подряд имитировать сбой отчёта

    async def find_order_by_name(self, name):
        found = self.orders.get(name)
        return dict(found) if found else None

    async def get_order_positions(self, order_id):
        return [{
            "quantity": 1.0,
            "assortment": {
                "id": "prod-1", "name": "Тестовый товар",
                "code": "T1", "uom": {"name": "шт"},
                "barcodes": [{"ean13": "4600000000001"}],
            },
        }]

    async def stock_by_slot(self, store_id, assortment_ids, chunk_size=50):
        if self.stock_fails > 0:
            self.stock_fails -= 1
            raise MoySkladError("МойСклад 503: сервис временно недоступен")
        return [{"assortmentId": "prod-1", "storeId": store_id,
                 "slotId": "slot-a", "stock": 10}]


class FakeBot:
    def __init__(self):
        self.documents = []

    async def send_document(self, chat_id, document, caption="", **kwargs):
        self.documents.append((chat_id, document, caption))


def order_row(oid, number, created, state="Упаковка"):
    return {"id": oid, "name": number, "created": created, "moment": created,
            "state": {"name": state}}


def make(tmp, client, bot, size=10, wait=60):
    settings = Settings(
        telegram_token="t", moysklad_token="m", notify_chat_id=-100500,
        notify_states=["Новый", "Упаковка"], batch_size=size, batch_max_wait=wait,
    )
    return PickingBatcher(
        bot, OrderService(client, "Склад адрес. хранение", clock=settings.clock()),
        settings, state_path=str(Path(tmp) / "batch.json"),
    )


def minutes_ago(clock, value):
    moment = clock.now_display() - timedelta(minutes=value)
    return moment.astimezone(clock.moysklad).strftime("%Y-%m-%d %H:%M:%S.000")


DEFAULT_CLOCK = Settings(telegram_token="t", moysklad_token="m").clock()


async def main():
    test_routing()
    test_sheet()

    # 1. Меньше порога и ждали недолго — лист не уходит
    with tempfile.TemporaryDirectory() as tmp:
        bot = FakeBot()
        client = BatchTestClient([])
        batcher = make(tmp, client, bot, size=3, wait=60)
        clock = batcher._clock
        batcher._load_state()

        client.orders["1001"] = order_row("o1", "1001", minutes_ago(clock, 5))
        client.orders["1002"] = order_row("o2", "1002", minutes_ago(clock, 3))
        batcher.enqueue(client.orders["1001"])
        batcher.enqueue(client.orders["1002"])
        assert await batcher.tick() is None
        assert bot.documents == []

        # Повторный enqueue того же заказа не плодит дублей в очереди
        batcher.enqueue(client.orders["1001"])
        assert len(batcher._queue) == 2, batcher._queue

        # 2. Набрался порог — уходит ровно size заказов, самые старые
        client.orders["1003"] = order_row("o3", "1003", minutes_ago(clock, 1))
        client.orders["1004"] = order_row("o4", "1004", minutes_ago(clock, 0))
        batcher.enqueue(client.orders["1003"])
        batcher.enqueue(client.orders["1004"])
        assert await batcher.tick() == 3
        assert len(bot.documents) == 1
        chat, document, caption = bot.documents[0]
        assert chat == -100500
        assert document.filename.startswith("Сборочный-лист-3-зак-")
        assert "Заказов: <b>3</b>" in caption, caption
        assert "1001, 1002, 1003" in caption and "1004" not in caption, caption
        assert [row["name"] for row in batcher._queue] == ["1004"], batcher._queue

        # 3. Повторный круг не отправляет то же самое (в очереди только 1004,
        #    порог не набран, ждать ещё не пора)
        assert await batcher.tick() is None

        # 4. Ожидание вышло — уходит остаток, даже если порог не набран.
        #    Ожидание отсчитывается от момента постановки В ОЧЕРЕДЬ
        #    (queued_at), а не от даты создания заказа в МойСклад — заказ
        #    мог встать в очередь позже, чем был создан.
        client.orders["1005"] = order_row("o5", "1005", minutes_ago(clock, 2))
        batcher.enqueue(client.orders["1005"])
        batcher._queue[0]["queued_at"] = minutes_ago(clock, 90)  # «1004» ждёт давно
        assert await batcher.tick() == 2
        assert len(bot.documents) == 2

        # 5. Отправленное и очередь переживают перезапуск
        saved = json.loads((Path(tmp) / "batch.json").read_text(encoding="utf-8"))
        assert set(saved["sent"]) == {"o1", "o2", "o3", "o4", "o5"}, saved
        assert saved["queue"] == [], saved

        again = make(tmp, client, FakeBot(), size=3, wait=60)
        again._load_state()
        assert await again.tick() is None
        assert again._bot.documents == []
        # Повторный enqueue уже отправленного заказа — no-op
        again.enqueue(client.orders["1001"])
        assert again._queue == []

    # 6. ГЛАВНЫЙ СЛУЧАЙ — гонка статусов. Заказ поставлен в очередь, пока
    #    ещё был в «Упаковка», но к моменту отправки в МойСклад уже
    #    «Передача». Раньше живой опрос батчера такой заказ просто не
    #    находил и он не отправлялся никогда. Теперь очередь не зависит от
    #    текущего статуса — карточка собирается заново на момент отправки.
    with tempfile.TemporaryDirectory() as tmp:
        client = BatchTestClient([order_row("o9", "2001", minutes_ago(DEFAULT_CLOCK, 90))])
        bot = FakeBot()
        batcher = make(tmp, client, bot, size=10, wait=0)
        batcher.enqueue(client.orders["2001"])  # видели в «Упаковка»
        client.orders["2001"]["state"] = {"name": "Передача"}  # успел уйти дальше
        assert await batcher.tick() == 1
        assert len(bot.documents) == 1, "заказ потерялся из-за смены статуса"
        assert "2001" in bot.documents[0][2]

    # 7. Отчёт по ячейкам недоступен — заказ остаётся в очереди, не пропадает
    with tempfile.TemporaryDirectory() as tmp:
        client = BatchTestClient([order_row("o10", "2002", minutes_ago(DEFAULT_CLOCK, 90))])
        client.stock_fails = 1
        bot = FakeBot()
        batcher = make(tmp, client, bot, size=10, wait=0)
        batcher.enqueue(client.orders["2002"])
        assert await batcher.tick() is None, "лист ушёл без ячеек"
        assert len(batcher._queue) == 1, "заказ выпал из очереди при сбое"
        assert await batcher.tick() == 1, "повтор не подхватил заказ"
        assert len(bot.documents) == 1

    # 8. Заказ пропал из МойСклад между постановкой в очередь и отправкой —
    #    остальные заказы в том же листе всё равно уходят
    with tempfile.TemporaryDirectory() as tmp:
        client = BatchTestClient([
            order_row("o11", "2003", minutes_ago(DEFAULT_CLOCK, 90)),
            order_row("o12", "2004", minutes_ago(DEFAULT_CLOCK, 90)),
        ])
        bot = FakeBot()
        batcher = make(tmp, client, bot, size=10, wait=0)
        batcher.enqueue(client.orders["2003"])
        batcher.enqueue(client.orders["2004"])
        del client.orders["2003"]  # удалён/архивирован
        assert await batcher.tick() == 2, "весь лист пропал из-за одного удалённого заказа"
        assert "2004" in bot.documents[0][2] and "2003" not in bot.documents[0][2]

    print("OK: очередь не зависит от текущего статуса, сбои не теряют заказы")


if __name__ == "__main__":
    asyncio.run(main())
