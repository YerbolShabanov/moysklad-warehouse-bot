"""Проверка выгрузки заказа в .xlsx: содержимое, порядок и настройки печати."""
import io
import sys
from datetime import datetime
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from openpyxl import load_workbook

from msbot.excel import order_filename, render_order_xlsx
from msbot.orders import OrderCard, OrderPosition, SlotStock


def position(name, barcode, row=None, place=None, level=None, qty=1, uom="шт"):
    p = OrderPosition(name=name, code="00007", barcode=barcode, quantity=qty, uom=uom)
    if row is not None:
        p.slots = [SlotStock(slot_name=f"Ряд {row} место {place} этаж {level}", stock=9)]
    return p


def main():
    card = OrderCard(
        id="x", number="1068634272", state="Выдан",
        moment=datetime(2026, 9, 9, 10, 9), sales_channel="ОПТ",
        agent='ООО "Ромашка"', store="Склад адрес. хранение",
        positions=[
            position("Автокресло SKYLER V2 (total black)", "4690624189487", 10, 11, 2),
            position("Ходунки SMILEY V2 (grey)", "0004690624101601", 2, 13, 1, qty=2),
            position("Коврик Soft Floor (watercolor)", "4690624101267", 2, 12, 1, qty=3),
            position("Позиция без ячейки", ""),
        ],
    )

    workbook = load_workbook(io.BytesIO(render_order_xlsx(card)))
    sheet = workbook.active

    assert sheet.title == "Сборка"
    assert order_filename(card) == "Заказ-1068634272.xlsx"

    # Шапка заказа над таблицей
    heading = [sheet.cell(row=r, column=1).value for r in range(1, 7)]
    assert heading[0] == "Заказ покупателя 1068634272"
    assert "Выдан" in heading[1] and "09.09.2026 10:09" in heading[2]
    assert "ОПТ" in heading[3]

    # Строка заголовков таблицы
    header_row = next(
        r for r in range(1, sheet.max_row + 1)
        if sheet.cell(row=r, column=1).value == "№"
    )
    head = [sheet.cell(row=header_row, column=c).value for c in range(1, sheet.max_column + 1)]
    assert head == ["№", "Наименование товара", "Штрихкод", "Кол-во", "Ячейка"], head

    # Нумерация идёт по маршруту сборки, подряд
    numbers = [sheet.cell(row=r, column=1).value
               for r in range(header_row + 1, sheet.max_row + 1)]
    assert numbers == [1, 2, 3, 4], numbers

    # Порядок обхода: Р2-М12 → Р2-М13 → Р10-М11 → без ячейки
    cells = [sheet.cell(row=r, column=5).value
             for r in range(header_row + 1, sheet.max_row + 1)]
    assert cells == ["Р2-М12-Э1", "Р2-М13-Э1", "Р10-М11-Э2", "—"], cells

    quantities = [sheet.cell(row=r, column=4).value
                  for r in range(header_row + 1, sheet.max_row + 1)]
    assert quantities == ["3", "2", "1", "1"], quantities

    # Штрихкод хранится текстом — иначе Excel срежет ведущий ноль
    barcode_cell = sheet.cell(row=header_row + 2, column=3)
    assert barcode_cell.value == "0004690624101601", barcode_cell.value
    assert barcode_cell.number_format == "@"

    # Настройки печати: подгон по ширине и повтор шапки на каждой странице
    assert sheet.page_setup.fitToWidth == 1
    assert sheet.print_title_rows == f"${header_row}:${header_row}"
    assert sheet.freeze_panes == f"A{header_row + 1}"

    # Колонка «Ед.» появляется только при неоднородных единицах
    card.positions[0].uom = "кг"
    head2 = load_workbook(io.BytesIO(render_order_xlsx(card))).active
    row2 = next(r for r in range(1, head2.max_row + 1)
                if head2.cell(row=r, column=1).value == "№")
    assert [head2.cell(row=row2, column=c).value for c in range(1, 7)][4] == "Ед."

    print("OK: xlsx собирается, порядок обхода сохранён, печать настроена")


if __name__ == "__main__":
    main()
